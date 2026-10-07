"""Excel ブックの読み込み（openpyxl）.

数式セルは Excel が保存した計算結果（キャッシュ）を読む. ツール（openpyxl など）で作ったブックは計算結果を持たないため、
そのままでは数式セルが空になる. `load_workbook_with_formula_fallback` は、その場合に数式の文字列（例: =SUM(B4:B5)）を
値として入れ、入れたセルの一覧を返す（呼び出し側が警告に使う）.
"""

from __future__ import annotations

from pathlib import Path
from typing import Any

INSTALL_HINT = (
    "Excel を読むには openpyxl が必要です。"
    'devtools-common を extra 付きで入れてください: pip install "devtools-common[excel] @ git+…"'
    "（または pip install openpyxl）"
)


def require_openpyxl() -> Any:
    """openpyxl を import して返す. 無ければ入れ方を書いた ImportError."""
    try:
        import openpyxl
    except ModuleNotFoundError as e:  # pragma: no cover - 環境依存
        raise ImportError(INSTALL_HINT) from e
    return openpyxl


def _formula_text(cell: Any) -> str | None:
    if getattr(cell, "data_type", None) != "f" or cell.value is None:
        return None
    v = cell.value
    if isinstance(v, str):
        return v
    text_attr = getattr(v, "text", None)  # 配列数式（ArrayFormula）
    return text_attr if isinstance(text_attr, str) else str(v)


def load_workbook_with_formula_fallback(path: str | Path) -> tuple[Any, list[str]]:
    """計算結果の値（data_only）で読んだブックと、計算結果が無く数式の文字列を入れたセル（「シート名!A1」）の一覧.

    計算結果が保存されていない数式セルには、数式の文字列を値として入れる（表示形式・結合などはそのまま）.
    ファイルが無ければ FileNotFoundError.
    """
    openpyxl = require_openpyxl()
    path = Path(path)
    if not path.is_file():
        raise FileNotFoundError(f"Excel が見つかりません: {path}")
    wb = openpyxl.load_workbook(path, data_only=True)
    formulas = openpyxl.load_workbook(path, read_only=True)
    missing: list[str] = []
    try:
        for fws in formulas.worksheets:
            if not hasattr(fws, "iter_rows") or fws.title not in wb.sheetnames:
                continue
            ws = wb[fws.title]
            for row in fws.iter_rows():
                for c in row:
                    f = _formula_text(c)
                    if f is None:
                        continue
                    target = ws.cell(c.row, c.column)
                    if target.value is None:
                        target.value = f
                        missing.append(f"{fws.title}!{target.coordinate}")
    finally:
        formulas.close()
    return wb, missing
